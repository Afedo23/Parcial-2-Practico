"""P06, P07, P08 y reglas de calidad de datos del importador."""
import json
import unittest

import openpyxl
from base import AppTestCase, archivo_para_importar, crear_excel
from catalogo import servicios
from catalogo.importador import ErrorImportacion, importar


class TestImportacion(AppTestCase):
    def test_P06_importar_archivo_original(self):
        resumen, obs = self.importar_catalogo()
        self.assertEqual(self.db.valor("SELECT COUNT(DISTINCT codigo) FROM servicio_n1"), 12)
        self.assertEqual(self.db.valor("SELECT COUNT(DISTINCT codigo) FROM servicio_n2"), 46)
        self.assertTrue(resumen["controles"]["ok"], resumen["controles"])
        self.assertEqual((resumen["n1"]["creados"], resumen["n2"]["creados"]), (12, 46))
        self.assertGreater(resumen["observaciones"], 0, "las incidencias deben quedar registradas")
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM import_observacion WHERE ejecucion_id = ?", (resumen["ejecucion_id"],)), len(obs))
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM servicio_n2 WHERE nivel1_id IS NULL"), 0)

    def test_P06_codigos_se_conservan_como_texto(self):
        self.importar_catalogo()
        for cod in ("SE.12.1", "SE.12.2", "SE.12.3"):
            r = self.db.uno("SELECT codigo, codigo_original FROM servicio_n2 WHERE codigo = ?", (cod,))
            self.assertIsNotNone(r, cod)
            self.assertEqual(r["codigo_original"], cod)

    def test_P07_repetir_importacion_no_duplica_y_es_trazable(self):
        r1, _ = self.importar_catalogo()
        antes = (self.db.valor("SELECT COUNT(*) FROM servicio_n1"), self.db.valor("SELECT COUNT(*) FROM servicio_n2"))
        ruta, _ = archivo_para_importar()
        r2, _ = importar(self.db, ruta)
        despues = (self.db.valor("SELECT COUNT(*) FROM servicio_n1"), self.db.valor("SELECT COUNT(*) FROM servicio_n2"))
        self.assertEqual(antes, despues)
        self.assertEqual((r2["n1"]["creados"], r2["n1"]["actualizados"], r2["n2"]["creados"], r2["n2"]["actualizados"]), (0, 0, 0, 0))
        self.assertEqual((r2["n1"]["omitidos"], r2["n2"]["omitidos"]), (12, 46))
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM import_ejecucion"), 2)
        self.assertEqual(r1["observaciones"], r2["observaciones"], "las incidencias se reportan en cada corrida")
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM (SELECT codigo FROM servicio_n2 GROUP BY codigo HAVING COUNT(*) > 1) d"), 0)

    def test_P08_conflicto_SE12_una_sola_entidad_con_evidencia(self):
        self.importar_catalogo()
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM servicio_n1 WHERE codigo = ?", ("SE.12",)), 1)
        n1 = self.db.uno("SELECT * FROM servicio_n1 WHERE codigo = ?", ("SE.12",))
        nombres = {"Suministrar Analitica", "Mantener Tableros de Control"}
        self.assertIn(n1["nombre"], nombres)
        alternos = json.loads(n1["nombres_alternos"])
        self.assertEqual(set(alternos) | {n1["nombre"]}, nombres, "ambos valores deben quedar como evidencia")
        self.assertEqual(n1["estado_revision"], "REVISION")
        tipos = {r["tipo"] for r in self.db.consultar("SELECT tipo FROM import_observacion WHERE codigo = ?", ("SE.12",))}
        self.assertIn("CONFLICTO_NOMBRE_N1", tipos)

    def test_P08_atributos_ausentes_se_conservan_como_desconocidos(self):
        self.importar_catalogo()
        for cod in ("SE.12.1", "SE.12.2", "SE.12.3"):
            r = self.db.uno("SELECT * FROM servicio_n2 WHERE codigo = ?", (cod,))
            self.assertEqual(r["estado_revision"], "REVISION", cod)
            if self.usa_fixture:
                for col in ("indicador_activo", "clase_id", "criticidad_id", "tipo_id", "metrica", "minimo", "maximo"):
                    self.assertIsNone(r[col], f"{cod}.{col} no debe inventarse")
            self.assertTrue(r["minimo"] is None or r["minimo"] != 0 or self.usa_fixture is False)

    def test_celdas_combinadas_y_filas_de_continuacion_con_fixture(self):
        self.importar_catalogo()
        if not self.usa_fixture:
            self.skipTest("Comprobación específica del archivo de prueba")
        s = self.db.uno("SELECT * FROM servicio_n2 WHERE codigo = ?", ("SE.01.01",))
        self.assertEqual(s["nombre"], "Servicio SE.01.01")
        self.assertEqual(s["descripcion"], "Descripción de SE.01.01")  # valor de la celda principal del rango combinado
        self.assertEqual(s["origen_rango"], "A5:L7")  # el servicio ocupa 3 filas físicas y es UNO solo
        n1 = self.db.valor("SELECT codigo FROM servicio_n1 WHERE id = ?", (s["nivel1_id"],))
        self.assertEqual(n1, "SE.01")
        huerfana = [o for o in self.db.consultar("SELECT * FROM import_observacion WHERE tipo = ?", ("FILA_SIN_CODIGO",))]
        self.assertEqual(len(huerfana), 1)
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM servicio_n2 WHERE descripcion LIKE ?", ("%Continuación suelta%",)), 0)
        self.assertEqual(self.db.valor("SELECT clase_id FROM servicio_n2 WHERE codigo = ?", ("SE.01.02",)) is not None, True)
        d = self.db.uno("SELECT valor_origen, valor_destino FROM mapeo_etiquetas WHERE campo = ?", ("clase",))
        self.assertEqual((d["valor_origen"], d["valor_destino"]), ("recurrente", "RECURRENTE"))  # corrección registrada en mapeo
        n12_3 = self.db.uno("SELECT nivel1_id FROM servicio_n2 WHERE codigo = ?", ("SE.12.3",))
        self.assertEqual(self.db.valor("SELECT codigo FROM servicio_n1 WHERE id = ?", (n12_3["nivel1_id"],)), "SE.12")

    def test_reimportar_conserva_asignaciones_y_baja_logica(self):
        self.importar_catalogo()
        srv = self.db.valor("SELECT id FROM servicio_n2 ORDER BY codigo LIMIT 1")
        servicios.asignar(self.db, srv, self.org["seccion"], self.consulta_id)
        servicios.cambiar_estado_n2(self.db, self.db.valor("SELECT id FROM servicio_n2 ORDER BY codigo DESC LIMIT 1"), False)
        ruta, _ = archivo_para_importar()
        importar(self.db, ruta)
        r = self.db.uno("SELECT seccion_id, usuario_id FROM servicio_n2 WHERE id = ?", (srv,))
        self.assertEqual((r["seccion_id"], r["usuario_id"]), (self.org["seccion"], self.consulta_id))
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM servicio_n2 WHERE activo = ?", (False,)), 1)

    def test_archivo_inexistente_o_estructura_invalida(self):
        with self.assertRaises(ErrorImportacion):
            importar(self.db, "/ruta/que/no/existe.xlsx")
        import tempfile, os

        tmp = os.path.join(tempfile.gettempdir(), "malo.xlsx")
        crear_excel(tmp)
        wb = openpyxl.load_workbook(tmp)
        wb.active["C4"] = "OTRA COSA"
        wb.save(tmp)
        with self.assertRaises(ErrorImportacion):
            importar(self.db, tmp)
        self.assertEqual(self.db.valor("SELECT COUNT(*) FROM servicio_n2"), 0, "no debe importarse nada")

    def test_importador_no_modifica_el_excel_original(self):
        import hashlib

        ruta, _ = archivo_para_importar()
        antes = hashlib.sha256(ruta.read_bytes()).hexdigest()
        importar(self.db, ruta)
        self.assertEqual(antes, hashlib.sha256(ruta.read_bytes()).hexdigest())

    def test_umbral_invertido_y_ausente_no_se_convierte_en_cero(self):
        import tempfile, os

        tmp = os.path.join(tempfile.gettempdir(), "umbral.xlsx")
        crear_excel(tmp)
        wb = openpyxl.load_workbook(tmp)
        ws = wb.active
        # SE.12.1 (fila 99): mínimo > máximo; SE.12.2 (fila 100): solo máximo
        ws["K99"], ws["L99"], ws["L100"] = 50, 10, 20
        wb.save(tmp)
        importar(self.db, tmp)
        a = self.db.uno("SELECT minimo, maximo FROM servicio_n2 WHERE codigo = ?", ("SE.12.1",))
        self.assertEqual((a["minimo"], a["maximo"]), (None, None))
        b = self.db.uno("SELECT minimo, maximo FROM servicio_n2 WHERE codigo = ?", ("SE.12.2",))
        self.assertIsNone(b["minimo"])
        self.assertEqual(b["maximo"], 20)
        tipos = {r["tipo"] for r in self.db.consultar("SELECT tipo FROM import_observacion")}
        self.assertIn("UMBRAL_INVERTIDO", tipos)


if __name__ == "__main__":
    unittest.main()
